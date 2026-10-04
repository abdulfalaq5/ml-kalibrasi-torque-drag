"""Parser dua format client vs angka di sel Excel (titik dicek manual pada fixture sintetis)."""

import openpyxl
import pytest
from app.parsers.workbook import parse_workbook, section_from_filename, well_type_from_path
from openpyxl import Workbook

from tests.conftest import FIXTURES

ROADMAP = FIXTURES / "P_CONTOH1_8.5in TnD Roadmap.xlsx"
WELLPLAN = FIXTURES / "P_CONTOH2_BHA2A_12.25in_TnD.xlsm"


def _find(rows, op, depth, ff=None, sheet=None):
    hits = [
        r
        for r in rows
        if r.operation == op
        and r.depth == depth
        and r.ff == ff
        and (sheet is None or r.sheet == sheet)
    ]
    assert len(hits) == 1, (op, depth, ff, hits)
    return hits[0]


def test_roadmap_format_a_five_checkpoints():
    pw = parse_workbook(ROADMAP, ROADMAP.name, f"J/CONTOH1/{ROADMAP.name}")
    assert pw.fmt == "roadmap" and pw.kinds == {"plan", "actual"}
    assert not pw.has_errors, pw.issues
    # 1-2: Drag baris 8 -> kedalaman 8500, Tripping In FF 0.10 = 184.8, Tripping Out FF 0.30 = '247.3' (teks)
    assert _find(pw.plan, "slack_off", 8500, 0.1).value == 184.8
    r = _find(pw.plan, "pick_up", 8500, 0.3)
    assert (r.value, r.unit) == (247.3, "klbf")
    # 3: Torque baris 8 -> 7500 ft, Rotating Off Bottom FF 0.10 = 3284.9 ft-lbf
    r = _find(pw.plan, "torque_off_bottom", 7500, 0.1)
    assert (r.value, r.unit) == (3284.9, "ft-lbf")
    # 4-5: T&D Actual Reading baris 5 -> 7593 ft, PU 224.4 Klbs, torque on bottom 11882 Lbs-ft
    assert _find(pw.actual, "pick_up", 7593).value == 224.4
    r = _find(pw.actual, "torque_on_bottom", 7593)
    assert (r.value, r.unit) == (11882, "ft-lbf")
    # sel kosong tidak menjadi titik
    assert not [x for x in pw.actual if x.depth == 7593 and x.operation == "slack_off"]
    # blok "Graph reference" (kurva + kalibrasi) diabaikan
    assert {r.ff for r in pw.plan if r.operation == "pick_up"} == {0.1, 0.3, 0.5}
    assert pw.meta["block_weight_klbf"] == 20
    assert pw.meta["section_meta_in"] == 8.5
    assert pw.meta["calibration_drag_klbf"]["pick_up"] == 10
    assert pw.meta["well_type_folder"] == "J" and pw.meta["well_folder"] == "CONTOH1"


def test_wellplan_format_b_checkpoints():
    pw = parse_workbook(WELLPLAN, WELLPLAN.name, f"Horizontal/CONTOH2/{WELLPLAN.name}")
    assert pw.fmt == "wellplan" and not pw.has_errors, pw.issues
    raw = openpyxl.load_workbook(WELLPLAN, read_only=True, data_only=True)
    tla = list(raw["Tripping Load Analysis"].iter_rows(values_only=True))
    row = tla[59]  # baris 60
    depth = float(row[1])
    # header: Bit Depth | 0.5..0.2 Trip IN | Rotate Off Bottom | 0.2..0.5 Trip Out ; satuan 1000 lbf
    assert _find(pw.plan, "slack_off", depth, 0.3).value == row[4]
    assert _find(pw.plan, "pick_up", depth, 0.5).value == row[10]
    assert _find(pw.plan, "rotating_weight", depth, None).value == row[6]
    obt = list(raw["Off Bottom Torque analysis"].iter_rows(values_only=True))[59]
    r = _find(pw.plan, "torque_off_bottom", float(obt[1]), 0.3)
    assert (r.value, r.unit) == (obt[3], "kft-lbf")
    # Drilling Data baris 6: 3479 | SO 79 | ROT 89 | PU 111 | RPM | off 4.2 | | on 9.1 (tanpa satuan)
    assert _find(pw.actual, "pick_up", 3479, sheet="Drilling Data").value == 111
    r = _find(pw.actual, "torque_on_bottom", 3479, sheet="Drilling Data")
    assert (r.value, r.unit) == (9.1, "kft-lbf")
    # Rotary Drill Buckling Outputs: Surface Torque = rencana torque on bottom (Base FF 0.4)
    assert _find(pw.plan, "torque_on_bottom", 3400, 0.4).value == 10.45
    assert pw.meta["block_weight_klbf"] == 21 and pw.meta["mud_weight_ppg"] == 10.0
    assert pw.meta["bha_components"] >= 4 and pw.meta["casing_shoe"] > 0
    assert len(pw.survey) > 50 and pw.survey_units["dls"] == "deg/100ft"


@pytest.mark.parametrize(
    "name,expected",
    [
        ("P_MINA25_0025_BHA2A_8.5in_TnD.xlsm", 8.5),
        ("P_LISE25_43 (S3-E4-32D)_12.25in TnD Roadmap.xlsx", 12.25),
        ("P_BNKO26_LQR_PX69A_8.50_in TnD Roadmap_TD.xlsx", 8.5),
        ("P_PETA25_0008HW_BHA1A_22inHS_TnD.xlsm", 22.0),
        ("RGU24-P08 (S1-E1-34A)_BDSI-19_17.25 HS TnD Roadmap.xlsx", 17.25),
        ("P_MINA25_0017HW_BHA2A_12.25in Motor825_10.375in Stab.xlsm", 12.25),
        ("P_KOTA25_0002HW_MSF 12C_BHA3A_6.125in_OrbitG2.xlsm", 6.125),
        ("tanpa_ukuran.xlsx", None),
    ],
)
def test_section_from_filename(name, expected):
    assert section_from_filename(name) == expected


def test_well_type_from_path():
    assert well_type_from_path("Horizontal/MINAS X/file.xlsm") == "Horizontal"
    assert well_type_from_path("S/AMPUH/file.xlsx") == "S"
    assert well_type_from_path("SUMUR/file.xlsx") is None


def _save(tmp_path, sheets: dict[str, list[list]], name="uji.xlsx"):
    wb = Workbook()
    wb.remove(wb.active)
    for sname, rows in sheets.items():
        ws = wb.create_sheet(sname)
        for r in rows:
            ws.append(r)
    p = tmp_path / name
    wb.save(p)
    return p


def test_unknown_format(tmp_path):
    pw = parse_workbook(_save(tmp_path, {"Lainnya": [["a", "b"]]}))
    assert pw.has_errors and "Unknown format" in pw.issues[0].message


def test_not_excel(tmp_path):
    p = tmp_path / "palsu.xlsx"
    p.write_bytes(b"bukan excel")
    assert parse_workbook(p).has_errors


def test_roadmap_without_header(tmp_path):
    p = _save(tmp_path, {"Drag": [["x"]], "Torque": [["y"]]})
    pw = parse_workbook(p)
    assert any("Run Measured Depth" in i.message for i in pw.issues)


def test_out_of_range_values_dropped(tmp_path):
    head = [
        "Run Measured Depth using:",
        "Tripping In using:",
        "Tripping Out using:",
        "Rotating Off Bottom using:",
    ]
    ff = ["open hole friction factor: 0.30"] * 4
    p = _save(
        tmp_path,
        {
            "Drag": [
                [],
                [],
                [],
                head,
                ff,
                ["(ft)", "(kip)", "(kip)", "(kip)"],
                [1000, 50, 60, 55],
                [2000, 70, 99999, 75],
            ],
            "Torque": [
                [],
                [],
                [],
                [],
                [
                    "Run Measured Depth using:",
                    "Rotating On Bottom using:",
                    "Rotating Off Bottom using:",
                ],
                ["open hole friction factor: 0.30"] * 3,
                ["(ft)", "(ft-lbf)", "(ft-lbf)"],
                [1000, 2000, 1000],
            ],
        },
    )
    pw = parse_workbook(p)
    assert any("outside the plausible range" in i.message for i in pw.issues)
    assert not any(r.value == 99999 for r in pw.plan)


@pytest.mark.parametrize(
    "md,inc,expected",
    [
        ([0, 1000, 2000, 3000], [0, 30, 45, 45], "J"),
        ([0, 1000, 2000, 3000, 4000], [0, 30, 40, 20, 10], "S"),
        ([0, 1000, 2000, 3000], [0, 40, 85, 90], "Horizontal"),
    ],
)
def test_classify_well_type(md, inc, expected):
    from app.services.classify import classify_well_type

    assert classify_well_type(md, inc)[0] == expected


def test_classify_section():
    from app.services.classify import classify_section

    assert classify_section(12.25) == 12.25
    assert classify_section(17.25) == 17.5
    assert classify_section(22) == 22.0
    assert classify_section(None) is None
