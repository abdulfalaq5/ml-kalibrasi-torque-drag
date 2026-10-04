"""Template unduhan: dibuat, diisi seperti pengguna, diunggah, diproses."""

import io

import openpyxl
from app.parsers.workbook import parse_workbook
from app.services.templates import build_template

MD = [3000, 4000, 5000, 6000, 7000]


def fill(kind: str, path, name="TPL-01", wtype="J", section="8.5", actual=True):
    wb = openpyxl.load_workbook(io.BytesIO(build_template(kind)))
    info = wb["Well Info"]
    info["B3"], info["B4"], info["B5"], info["B6"], info["B7"] = name, section, wtype, 20, 2800
    drag = wb["Drag"]  # data mulai baris 7; blok FF di kolom A, F, K
    for b, c0 in enumerate((1, 6, 11)):
        for i, d in enumerate(MD):
            r = 7 + i
            pu = 100 + d / 50 + b * 8
            so = 90 + d / 80 - b * 6
            for j, v in enumerate((d, so, pu, 95 + d / 60)):
                drag.cell(r, c0 + j, v)
    tq = wb["Torque"]  # data mulai baris 8; blok di kolom A, E, I
    for b, c0 in enumerate((1, 5, 9)):
        for i, d in enumerate(MD):
            for j, v in enumerate((d, 3000 + d * (0.6 + b * 0.3), 1000 + d * (0.5 + b * 0.3))):
                tq.cell(8 + i, c0 + j, v)
    if actual:
        act = wb["T&D Actual Reading"]  # data mulai baris 5
        for i, d in enumerate(range(3100, 6900, 300)):
            for j, v in enumerate(
                (
                    d,
                    100 + d / 50 + 9,
                    90 + d / 80 - 5,
                    95 + d / 60 + 2,
                    1000 + d * 0.8,
                    3000 + d * 0.9,
                )
            ):
                act.cell(5 + i, 1 + j, v)
    sv = wb["Survey"]  # data mulai baris 5
    for i, (md, inc) in enumerate([(0, 0), (2000, 10), (4000, 28), (7000, 30)]):
        for j, v in enumerate((md, inc, 140, 1.0)):
            sv.cell(5 + i, 1 + j, v)
    wb.save(path)
    return path


def test_template_sheets():
    for kind, has_actual in (("training", True), ("monitoring", False)):
        wb = openpyxl.load_workbook(io.BytesIO(build_template(kind)))
        names = set(wb.sheetnames)
        assert {"Instructions", "Well Info", "Drag", "Torque", "Survey", "Example Drag"} <= names
        assert ("T&D Actual Reading" in names) == has_actual


def test_empty_template_is_rejected_with_reason(tmp_path):
    p = tmp_path / "kosong.xlsx"
    p.write_bytes(build_template("monitoring"))
    pw = parse_workbook(p, p.name)
    assert pw.has_errors and any("WellPlan" in i.message for i in pw.issues)


def test_filled_template_parses(tmp_path):
    p = fill("training", tmp_path / "isi.xlsx")
    pw = parse_workbook(p, p.name)
    assert not pw.has_errors, pw.issues
    assert pw.meta["well_name"] == "TPL-01" and pw.meta["well_type_template"] == "J"
    assert pw.meta["section_template"] == 8.5 and pw.meta["block_weight_klbf"] == 20
    assert {r.ff for r in pw.plan if r.operation == "pick_up"} == {0.1, 0.3, 0.5}
    assert {r.operation for r in pw.plan} == {
        "pick_up",
        "slack_off",
        "rotating_weight",
        "torque_on_bottom",
        "torque_off_bottom",
    }
    assert len({r.depth for r in pw.actual}) >= 12
    assert len(pw.survey) == 4
    # sheet Contoh tidak ikut terbaca
    assert all(r.depth in MD for r in pw.plan)


def test_upload_filled_template(auth_client, tmp_path):
    p = fill("training", tmp_path / "upload.xlsx", name="TPL-UP", wtype="S")
    with p.open("rb") as fh:
        j = auth_client.post(
            "/api/files",
            files={"file": (p.name, fh)},
            data={"purpose": "training", "section_in": "8.5", "well_type": "S"},
        ).json()
    assert j["status"] in ("ok", "warning"), j
    assert j["purpose"] == "training"
    assert j["well_name"] == "TPL-UP" and j["well_type"] == "S" and j["section_in"] == 8.5
    assert j["quality"]["status"] in ("A", "B", "C")
    r = auth_client.get("/api/templates/monitoring.xlsx")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    r = auth_client.get("/api/templates/sumur-baru.xlsx")  # tautan lama tetap jalan
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert auth_client.get("/api/templates/lainnya.xlsx").status_code == 404
