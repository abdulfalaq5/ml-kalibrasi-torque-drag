"""Perbaikan dari feedback client 01: kalibrasi DD, label seri standar, pemisahan
Training / Monitoring, template bahasa Inggris (template lama tetap terbaca)."""

import io

import openpyxl
import pytest
from app.db.models import PlanResult, Well
from app.parsers.workbook import parse_workbook
from app.services import units
from app.services.dataset import build_dataset
from app.services.profile import well_profile
from app.services.quality import eligible_well_ids
from app.services.templates import build_template
from openpyxl import Workbook

from tests.test_templates import fill

DRAG_HEAD = [
    "Run Measured Depth using:",
    "Tripping In using:",
    "Tripping Out using:",
    "Rotating Off Bottom using:",
]
TQ_HEAD = ["Run Measured Depth using:", "Rotating On Bottom using:", "Rotating Off Bottom using:"]


def _roadmap(tmp_path, torque_off_label="Calibrate Off Bot Torque"):
    wb = Workbook()
    wb.remove(wb.active)
    drag = wb.create_sheet("Drag")
    for r in (
        ["Calibrate", None, "PICK UP", "SLACK OFF", "ROTATE"],
        [None, None, 12, -3, 4.5],
        [],
        DRAG_HEAD,
        ["open hole friction factor: 0.30"] * 4,
        ["(ft)", "(kip)", "(kip)", "(kip)"],
        [1000, 50, 60, 55],
        [2000, 70, 90, 75],
    ):
        drag.append(r)
    tq = wb.create_sheet("Torque")
    for r in (
        ["WellPlan Result"],
        ["Calibrate On Bot Torque", None, 1500],
        [torque_off_label, None, 1800],
        [],
        TQ_HEAD,
        ["open hole friction factor: 0.30"] * 3,
        ["(ft)", "(ft-lbf)", "(ft-lbf)"],
        [1000, 2000, 1000],
        [2000, 3000, 1500],
    ):
        tq.append(r)
    p = tmp_path / "cal.xlsx"
    wb.save(p)
    return p


def test_drag_and_torque_calibration_parsed(tmp_path):
    pw = parse_workbook(_roadmap(tmp_path))
    assert pw.meta["calibration_drag_klbf"] == {"pick_up": 12, "slack_off": -3, "rotate": 4.5}
    assert pw.meta["calibration_torque_ftlbf"] == {"on_bottom": 1500, "off_bottom": 1800}


def test_torque_off_calibration_read_by_position(tmp_path):
    # di beberapa file client label baris ke-2 tertimpa angka; nilai tetap di kolom C
    pw = parse_workbook(_roadmap(tmp_path, torque_off_label=7))
    assert pw.meta["calibration_torque_ftlbf"] == {"on_bottom": 1500, "off_bottom": 1800}


def test_template_has_calibration_rows_and_english_sheets():
    wb = openpyxl.load_workbook(io.BytesIO(build_template("training")))
    assert wb["Drag"]["A1"].value == "Calibrate" and wb["Drag"]["C1"].value == "PICK UP"
    assert wb["Torque"]["A2"].value == "Calibrate On Bot Torque"
    assert wb["Well Info"]["A3"].value == "Well name"


def test_legacy_indonesian_template_still_parses(tmp_path):
    p = fill("training", tmp_path / "lama.xlsx")
    wb = openpyxl.load_workbook(p)
    ws = wb["Well Info"]
    ws.title = "Info Sumur"
    for cell, label in (("A3", "Nama sumur"), ("A4", "Section (inci)"), ("A5", "Tipe sumur")):
        ws[cell] = label
    wb.save(p)
    pw = parse_workbook(p, p.name)
    assert not pw.has_errors, pw.issues
    assert pw.meta["well_name"] == "TPL-01" and pw.meta["section_template"] == 8.5
    assert pw.meta["well_type_template"] == "J"


@pytest.fixture
def cal_well(db):
    w = Well(
        name="CAL-01",
        section_in=8.5,
        well_type="J",
        meta={
            "calibration_drag_klbf": {"pick_up": 10, "slack_off": -5, "rotate": 2},
            "calibration_torque_ftlbf": {"on_bottom": 1000, "off_bottom": 500},
        },
    )
    db.add(w)
    db.flush()
    for ff in (0.3, 0.5):
        for op, unit, base in (
            ("pick_up", "klbf", 100),
            ("slack_off", "klbf", 80),
            ("rotating_weight", "klbf", 90),
            ("torque_on_bottom", "ft-lbf", 5000),
        ):
            for d in (1000, 2000):
                v = base + d / 100 + ff * 10
                db.add(
                    PlanResult(
                        well_id=w.id,
                        operation=op,
                        ff=ff,
                        source_sheet="Drag",
                        depth_m=units.to_si(d, "ft"),
                        value=v,
                        unit=unit,
                        value_si=units.to_si(v, unit),
                    )
                )
    db.commit()
    yield w
    db.delete(w)
    db.commit()


def test_profile_calibrated_curves_and_series_names(db, cal_well):
    raw = well_profile(db, cal_well, "imperial", calibration="raw")
    cal = well_profile(db, cal_well, "imperial")  # default: calibrated bila ada offset
    assert cal["calibration"] == {"available": True, "mode": "calibrated"}
    pu_raw, pu_cal = raw["operations"]["pick_up"], cal["operations"]["pick_up"]
    assert [s["name"] for s in pu_cal["wellplan"]] == ["PU - OHFF : 0.3", "PU - OHFF : 0.5"]
    for a, b in zip(pu_raw["wellplan"][0]["value"], pu_cal["wellplan"][0]["value"], strict=True):
        assert b - a == pytest.approx(10.0, abs=1e-3)
    so = cal["operations"]["slack_off"]["wellplan"]
    assert so[0]["name"] == "SO - OHFF : 0.3"
    assert so[0]["value"][0] - raw["operations"]["slack_off"]["wellplan"][0]["value"][0] == (
        pytest.approx(-5.0, abs=1e-3)
    )
    rot = cal["operations"]["rotating_weight"]["wellplan"]
    assert [s["name"] for s in rot] == ["ROT"]  # satu kurva, tanpa OHFF
    ton = cal["operations"]["torque_on_bottom"]
    assert ton["wellplan"][1]["name"] == "Torque On Bottom - OHFF : 0.5"
    assert ton["calibration_offset"] == pytest.approx(1000.0)


def test_upload_requires_purpose_section_and_type(auth_client, tmp_path):
    p = fill("training", tmp_path / "x.xlsx", name="REQ-01")
    with p.open("rb") as fh:
        r = auth_client.post("/api/files", files={"file": (p.name, fh)})
    assert r.status_code == 422
    with p.open("rb") as fh:
        r = auth_client.post(
            "/api/files",
            files={"file": (p.name, fh)},
            data={"purpose": "training", "section_in": "9.1", "well_type": "J"},
        )
    assert r.status_code == 400 and "section" in r.json()["detail"].lower()


def test_monitoring_upload_never_enters_training(auth_client, db, tmp_path):
    """File monitoring dengan data aktual lengkap (kualitas bisa A) tidak pernah masuk dataset."""
    p = fill("training", tmp_path / "mon.xlsx", name="SEP-01")
    form = {"section_in": "8.5", "well_type": "J"}
    with p.open("rb") as fh:
        mon = auth_client.post(
            "/api/files", files={"file": (p.name, fh)}, data={**form, "purpose": "monitoring"}
        ).json()
    assert mon["status"] in ("ok", "warning") and mon["purpose"] == "monitoring"
    mon_id = mon["well_id"]

    ids, _ = eligible_well_ids(db)
    assert mon_id not in ids
    ds, _ = build_dataset(db, ids)
    assert ds.empty or mon_id not in set(ds.well_id)
    assert all(
        w["purpose"] == "monitoring"
        for w in auth_client.get("/api/wells?purpose=monitoring").json()
    )
    assert mon_id not in {w["id"] for w in auth_client.get("/api/wells?purpose=training").json()}

    # file yang sama sebagai training = sumur terpisah (tidak tercampur)
    with p.open("rb") as fh:
        tr = auth_client.post(
            "/api/files", files={"file": (p.name, fh)}, data={**form, "purpose": "training"}
        ).json()
    assert tr["well_id"] != mon_id and tr["purpose"] == "training"
    ids2, _ = eligible_well_ids(db)
    assert mon_id not in ids2

    # promote hanya untuk sumur monitoring
    assert auth_client.post(f"/api/wells/{tr['well_id']}/promote").status_code == 400
    for wid in (mon_id, tr["well_id"]):
        assert auth_client.delete(f"/api/wells/{wid}").status_code == 200


def test_forecast_backtest_bias_correction_removes_constant_offset():
    """Offset konstan 30 klbf per sumur: T&D mentah gagal toleransi, T&D + bias lolos semua."""
    import pandas as pd
    from app.services.backtest import forecast_backtest
    from app.services.metrics import tolerance_si, within_frac

    rows = []
    for wid in (1, 2):
        for k in range(40):
            depth = units.to_si(3000 + 50 * k, "ft")
            wp = units.to_si(100 + k, "klbf")
            rows.append(
                {
                    "well_id": wid,
                    "depth_m": depth,
                    "wp_base": wp,
                    "target": wp + units.to_si(30, "klbf"),
                    "ml_oof": wp + units.to_si(25, "klbf"),
                }
            )
    d = pd.DataFrame(rows)
    bt = forecast_backtest(d, "pick_up")
    assert set(bt["horizons"]) == {"300", "600", "1000"}
    h = bt["horizons"]["300"]
    assert h["T&D model"]["within"] == 0.0
    assert h["T&D model + bias"]["within"] == 1.0 and h["ML + bias"]["within"] == 1.0
    assert h["ML"]["within"] == 1.0  # 5 klbf < 8 klbf
    assert within_frac([0, 0], [1, 100], tolerance_si("pick_up")) == 0.5
