"""Tanpa login, semua endpoint ditolak (kecuali health dan login)."""

import re

from app.db.session import SessionLocal  # noqa: E402
from fastapi.routing import APIRoute
from sqlalchemy import delete  # noqa: E402

PUBLIC = {"/api/health", "/api/auth/login"}


def _concrete(path: str) -> str:
    return re.sub(r"\{[^}]+\}", "1", path).replace("1:path", "x")


def test_every_api_route_requires_session(client):
    from app.main import app

    checked = 0
    for route in app.routes:
        if not isinstance(route, APIRoute) or not route.path.startswith("/api"):
            continue
        if route.path in PUBLIC:
            continue
        for method in route.methods:
            r = client.request(method, _concrete(route.path))
            assert r.status_code == 401, (method, route.path, r.status_code)
            checked += 1
    assert checked >= 15
    # rute yang tidak ada juga 401 (tidak membocorkan daftar endpoint)
    assert client.get("/api/tidak-ada").status_code == 401


def test_health_public(client):
    r = client.get("/api/health")
    assert r.status_code == 200
    assert r.json()["database"] == "OK"


def test_docs_disabled_in_production(monkeypatch):
    from app.core.config import get_settings
    from app.main import create_app

    monkeypatch.setenv("APP_ENV", "production")
    get_settings.cache_clear()
    try:
        app = create_app()
        paths = {getattr(r, "path", None) for r in app.routes}
        assert "/docs" not in paths and "/openapi.json" not in paths
    finally:
        monkeypatch.setenv("APP_ENV", "test")
        get_settings.cache_clear()


def test_login_logout_me(auth_client):
    assert auth_client.get("/api/auth/me").json() == {"username": "admin", "role": "admin"}
    assert auth_client.get("/api/wells").status_code == 200
    assert auth_client.post("/api/auth/logout").status_code == 200
    assert auth_client.get("/api/wells").status_code == 401


def test_session_cookie_flags(client):
    r = client.post(
        "/api/auth/login", json={"username": "admin", "password": "password-uji-panjang"}
    )
    cookie = r.headers["set-cookie"].lower()
    assert "httponly" in cookie and "samesite=lax" in cookie
    assert "max-age=28800" in cookie


def test_lockout_after_five_failures(client, db):
    from app.db.models import LoginAttempt
    from sqlalchemy import delete

    db.execute(delete(LoginAttempt))
    db.commit()
    for _ in range(5):
        r = client.post("/api/auth/login", json={"username": "admin", "password": "salah"})
        assert r.status_code == 401
    # percobaan keenam terkunci, bahkan dengan password benar
    r = client.post(
        "/api/auth/login", json={"username": "admin", "password": "password-uji-panjang"}
    )
    assert r.status_code == 429
    db.execute(delete(LoginAttempt))
    db.commit()


def test_password_is_hashed(db):
    from app.db.models import AdminUser
    from sqlalchemy import select

    user = db.scalar(select(AdminUser))
    assert user.password_hash.startswith("$argon2")
    assert "password-uji-panjang" not in user.password_hash


def test_guest_only_monitoring_and_dashboard(auth_client, tmp_path):
    """Akun guest (K-45): hanya rute Monitoring/Dashboard dan sumur monitoring miliknya (K-47);
    isi tampilannya sama dengan admin, termasuk WellPlan, Excel, dan PDF (K-48)."""
    from app.core.security import GUEST_ROUTES, hash_password
    from app.db.models import AdminUser, LoginAttempt
    from app.main import app
    from fastapi.testclient import TestClient
    from sqlalchemy import select

    from tests.test_templates import fill

    with SessionLocal() as s:
        if s.scalar(select(AdminUser).where(AdminUser.username == "tamu-uji")) is None:
            s.add(
                AdminUser(username="tamu-uji", password_hash=hash_password("x" * 12), role="guest")
            )
        s.execute(delete(LoginAttempt))
        s.commit()
        guest_id = s.scalar(select(AdminUser.id).where(AdminUser.username == "tamu-uji"))

    form = {"section_in": "8.5", "well_type": "J"}
    ids = {}
    for purpose in ("training", "monitoring"):
        p = fill("training", tmp_path / f"{purpose}.xlsx", name=f"GUEST-{purpose}")
        # admin mengunggah sumur monitoring atas nama guest (K-47)
        owner = {"owner_id": str(guest_id)} if purpose == "monitoring" else {}
        with p.open("rb") as fh:
            r = auth_client.post(
                "/api/files",
                files={"file": (p.name, fh)},
                data={**form, "purpose": purpose, **owner},
            ).json()
        ids[purpose] = (r["well_id"], r["id"])
    tw, tf = ids["training"]
    mw, mf = ids["monitoring"]

    with TestClient(app) as g:
        r = g.post("/api/auth/login", json={"username": "tamu-uji", "password": "x" * 12})
        assert r.status_code == 200 and r.json()["role"] == "guest"
        assert g.get("/api/auth/me").json()["role"] == "guest"

        # semua rute di luar daftar guest ditolak (403), tanpa login tetap 401
        for route in app.routes:
            if not isinstance(route, APIRoute) or not route.path.startswith("/api"):
                continue
            if route.path in PUBLIC:
                continue
            for method in route.methods:
                if (method, route.path) in GUEST_ROUTES:
                    continue
                r = g.request(method, _concrete(route.path))
                assert r.status_code == 403, (method, route.path, r.status_code)

        # hanya sumur dan file monitoring
        assert {w["purpose"] for w in g.get("/api/wells").json()} == {"monitoring"}
        assert {w["purpose"] for w in g.get("/api/wells?purpose=training").json()} == {"monitoring"}
        assert {f["purpose"] for f in g.get("/api/files").json()} == {"monitoring"}
        assert g.get(f"/api/wells/{tw}/profile").status_code == 404
        assert g.delete(f"/api/wells/{tw}").status_code == 404
        assert g.get(f"/api/files/{tf}/download").status_code == 404
        assert g.get(f"/api/files/{mf}/download").status_code == 200
        assert g.get("/api/templates/training.xlsx").status_code == 403
        assert g.get(f"/api/wells/{tw}/export.xlsx").status_code == 404
        assert g.get(f"/api/wells/{tw}/report.pdf").status_code == 404
        assert g.get("/api/templates/monitoring.xlsx").status_code == 200
        p = fill("training", tmp_path / "g.xlsx", name="GUEST-try")
        with p.open("rb") as fh:
            r = g.post(
                "/api/files", files={"file": (p.name, fh)}, data={**form, "purpose": "training"}
            )
        assert r.status_code == 403

        # profil, Excel, dan PDF sama dengan yang dilihat admin (K-48)
        pr = g.get(f"/api/wells/{mw}/profile")
        assert pr.status_code == 200
        assert pr.json() == auth_client.get(f"/api/wells/{mw}/profile").json()
        ops = pr.json()["operations"].values()
        assert any(o["wellplan"] for o in ops) and all(o["actual"]["depth"] for o in ops)

        import io

        import openpyxl

        def sheets(c):
            x = c.get(f"/api/wells/{mw}/export.xlsx")
            assert x.status_code == 200, x.text
            wb = openpyxl.load_workbook(io.BytesIO(x.content))
            cells = {str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value}
            return wb.sheetnames, cells

        names, cells = sheets(g)
        assert (names, cells) == sheets(auth_client)
        assert "ML PREDICTION" in cells and any("MODELLED" in c for c in cells)
        pdf = g.get(f"/api/wells/{mw}/report.pdf")
        assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
        fx = g.post(f"/api/wells/{mw}/forecast.xlsx", json={"distance_ft": 300})
        # rute terbuka untuk guest (di tes ini belum ada model aktif -> 400, bukan 403)
        assert fx.status_code in (200, 400), fx.text

    # ekspor prediction dari respons yang sudah disaring untuk guest tetap bisa dibuat
    from app.services.forecast import export_forecast
    from app.services.guest_view import forecast_for_guest

    op = {
        "label": "Pick up",
        "unit": "klbf",
        "depth": [3000.0, 3100.0],
        "ml": [100.0, 102.0],
        "p10": [95.0, 97.0],
        "p90": [105.0, 107.0],
        "ml_corrected": None,
        "bias": None,
        "wellplan": [{"ff": 0.3, "name": "PU - OHFF : 0.3", "value": [99.0, 101.0]}],
        "change": 2.0,
        "tolerance": "8 klbf",
        "backtest": None,
        "actual_check": {
            "n": 2,
            "td_within": 0.5,
            "ml_within": 1.0,
            "ml_mean_abs": 1.0,
            "unit": "klbf",
        },
        "explanation": {
            "method": "SHAP",
            "drivers": [{"feature": "wp_base", "label": "T&D Model", "delta": 2.0}],
            "plan_changes": {"inclination_deg": [10, 12]},
            "limit_crossings": [],
            "sentence": "Pick up: from 3,000 to 3,100 ft the ML prediction is expected to rise from "
            "100.0 to 102.0 klbf (+2.0). Main drivers: T&D Model (+2.0).",
        },
    }
    fc = {
        "well": {
            "id": 1,
            "name": "X",
            "section_in": 8.5,
            "well_type": "J",
            "purpose": "monitoring",
        },
        "model_id": 1,
        "unit_system": "imperial",
        "depth_unit": "ft",
        "start_depth": 3000.0,
        "end_depth": 3100.0,
        "last_actual_depth": 3000.0,
        "distance_ft": 100,
        "bias_correction": False,
        "calibration": "raw",
        "warnings": [],
        "summary": "",
        "operations": {"pick_up": op},
    }
    g_fc = forecast_for_guest(fc)
    assert "T&D" not in g_fc["operations"]["pick_up"]["explanation"]["sentence"]
    wb = openpyxl.load_workbook(io.BytesIO(export_forecast(g_fc)))
    cells = {str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value}
    assert not any("T&D" in c or "OHFF" in c or "WellPlan" in c for c in cells)

    for wid in (tw, mw):
        auth_client.delete(f"/api/wells/{wid}")
    _delete_orphans()


def _delete_orphans() -> None:
    """SQLite uji tidak menegakkan foreign key: hapus baris yatim agar ID sumur yang dipakai
    ulang oleh tes berikutnya tidak mewarisi data aktual / rencana sumur yang dihapus."""
    from app.db.models import (
        ActualReading,
        PlanResult,
        Prediction,
        PredictionPoint,
        Survey,
        Well,
        WellQuality,
    )
    from sqlalchemy import select

    with SessionLocal() as s:
        for model in (ActualReading, PlanResult, Survey, Prediction, WellQuality):
            s.execute(delete(model).where(model.well_id.not_in(select(Well.id))))
        s.execute(
            delete(PredictionPoint).where(
                PredictionPoint.prediction_id.not_in(select(Prediction.id))
            )
        )
        s.commit()


def test_user_management_and_well_ownership(auth_client, tmp_path):
    """Admin CRUD akun (K-47); guest hanya melihat sumur monitoring miliknya, admin melihat semua."""
    from app.main import app
    from fastapi.testclient import TestClient

    from tests.test_templates import fill

    # validasi
    assert (
        auth_client.post(
            "/api/users", json={"username": "u-a", "password": "pendek", "role": "guest"}
        ).status_code
        == 400
    )
    assert (
        auth_client.post(
            "/api/users", json={"username": "u-a", "password": "x" * 12, "role": "boss"}
        ).status_code
        == 400
    )

    users = {}
    for name in ("pengguna-a", "pengguna-b"):
        r = auth_client.post(
            "/api/users", json={"username": name, "password": "rahasia-" + name, "role": "guest"}
        )
        assert r.status_code == 201, r.text
        users[name] = r.json()["id"]
    assert (
        auth_client.post(
            "/api/users", json={"username": "PENGGUNA-A", "password": "x" * 12}
        ).status_code
        == 409
    )
    listed = {u["username"]: u for u in auth_client.get("/api/users").json()}
    assert listed["pengguna-a"]["role"] == "guest" and "password_hash" not in listed["pengguna-a"]

    # admin tidak bisa menghapus / menurunkan dirinya sendiri (admin satu-satunya)
    admin_id = listed["admin"]["id"]
    assert auth_client.delete(f"/api/users/{admin_id}").status_code == 400
    assert auth_client.patch(f"/api/users/{admin_id}", json={"role": "guest"}).status_code == 400

    clients = {}
    for name in users:
        c = TestClient(app)
        c.__enter__()
        r = c.post("/api/auth/login", json={"username": name, "password": "rahasia-" + name})
        assert r.status_code == 200, r.text
        clients[name] = c
    a, b = clients["pengguna-a"], clients["pengguna-b"]
    try:
        assert a.get("/api/users").status_code == 403

        form = {"section_in": "8.5", "well_type": "J", "purpose": "monitoring"}
        p = fill("training", tmp_path / "a.xlsx", name="OWN-A")
        with p.open("rb") as fh:
            ra = a.post("/api/files", files={"file": (p.name, fh)}, data=form).json()
        assert ra["status"] in ("ok", "warning"), ra
        wa, fa = ra["well_id"], ra["id"]
        # guest tidak boleh memilih pemilik
        with p.open("rb") as fh:
            r = a.post(
                "/api/files",
                files={"file": (p.name, fh)},
                data={**form, "owner_id": str(users["pengguna-b"])},
            )
        assert r.status_code == 400

        # B tidak melihat sumur / file A
        assert [w["id"] for w in a.get("/api/wells").json()] == [wa]
        assert b.get("/api/wells").json() == []
        assert b.get("/api/files").json() == []
        assert b.get(f"/api/wells/{wa}/profile").status_code == 404
        assert b.delete(f"/api/wells/{wa}").status_code == 404
        assert b.get(f"/api/files/{fa}/download").status_code == 404
        assert a.patch(f"/api/wells/{wa}", json={"owner_id": None}).status_code == 403

        # K-54: B mengunggah file dengan nama sumur yang sama -> diterima sebagai sumur baru milik B
        # dengan stempel tanggal-waktu; sumur A tidak berubah
        with p.open("rb") as fh:
            rb = b.post("/api/files", files={"file": (p.name, fh)}, data=form).json()
        assert rb["status"] == "warning" and rb["well_id"] not in (None, wa), rb
        assert any("another user" in i["message"] for i in rb["issues"])
        wb_ = b.get("/api/wells").json()
        assert [w["id"] for w in wb_] == [rb["well_id"]]
        assert re.fullmatch(r"OWN-A \d{4}-\d\d-\d\d \d\d:\d\d:\d\d", wb_[0]["name"])
        assert [w["id"] for w in a.get("/api/wells").json()] == [wa]
        assert [f["id"] for f in b.get("/api/files").json()] == [rb["id"]]
        # A mengunggah ulang file yang sama -> tetap diproses, sumur baru berstempel
        with p.open("rb") as fh:
            ra2 = a.post("/api/files", files={"file": (p.name, fh)}, data=form).json()
        assert ra2["id"] != fa and ra2["well_id"] not in (None, wa, rb["well_id"]), ra2
        assert any("same file" in i["message"] for i in ra2["issues"])
        assert auth_client.delete(f"/api/wells/{ra2['well_id']}").status_code == 200
        assert auth_client.delete(f"/api/wells/{rb['well_id']}").status_code == 200

        # admin melihat semua, termasuk pemiliknya, dan bisa memindahkan pemilik
        mine = {w["id"]: w for w in auth_client.get("/api/wells?purpose=monitoring").json()}
        assert mine[wa]["owner"] == "pengguna-a"
        r = auth_client.patch(f"/api/wells/{wa}", json={"owner_id": users["pengguna-b"]})
        assert r.status_code == 200 and r.json()["owner"] == "pengguna-b"
        assert a.get("/api/wells").json() == []
        assert [w["id"] for w in b.get("/api/wells").json()] == [wa]

        # ganti password -> password lama tidak berlaku
        uid = users["pengguna-a"]
        assert (
            auth_client.patch(
                f"/api/users/{uid}", json={"password": "password-baru-a1"}
            ).status_code
            == 200
        )
        r = a.post(
            "/api/auth/login", json={"username": "pengguna-a", "password": "rahasia-pengguna-a"}
        )
        assert r.status_code == 401

        # hapus akun -> sesi langsung tidak berlaku, sumurnya tetap ada tanpa pemilik
        assert auth_client.delete(f"/api/users/{users['pengguna-b']}").status_code == 200
        assert b.get("/api/wells").status_code == 401
        assert auth_client.get(f"/api/wells/{wa}").json()["owner_id"] is None
    finally:
        for c in clients.values():
            c.__exit__(None, None, None)
        auth_client.delete(f"/api/wells/{wa}")
        auth_client.delete(f"/api/users/{users['pengguna-a']}")
        _delete_orphans()
        with SessionLocal() as s:
            from app.db.models import LoginAttempt

            s.execute(delete(LoginAttempt))
            s.commit()
