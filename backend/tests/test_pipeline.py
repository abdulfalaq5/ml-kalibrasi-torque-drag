"""Alur penuh paket 6 minggu dengan data sintetis berstruktur folder seperti data client:
pindai folder -> gerbang kualitas -> tinjauan -> bekukan dataset -> latih -> blind test ->
prediksi -> batas aman -> ekspor Excel/PDF -> evaluasi prediksi setelah data aktual masuk."""

import io
import os
import subprocess
import sys
import time
import zipfile

import openpyxl
import pytest
from app.db.models import MLModel, Prediction, Well
from app.services.training import _compare_and_activate, train
from make_sample_data import build_well, section_spec, well_specs, write_wellplan_file
from sqlalchemy import select

from tests.conftest import ROOT

SEED = 11


@pytest.fixture(scope="module")
def inbox(tmp_dir):
    path = tmp_dir / "inbox"
    subprocess.run(
        [
            sys.executable,
            str(ROOT / "scripts" / "make_sample_data.py"),
            "--out",
            str(path),
            "--wells",
            "9",
            "--new-wells",
            "1",
            "--seed",
            str(SEED),
        ],
        check=True,
        capture_output=True,
    )
    old = time.time() - 300
    for f in path.rglob("*.xls*"):
        os.utime(f, (old, old))
    return path


def test_upload_rejects_non_excel(auth_client, tmp_path):
    p = tmp_path / "x.csv"
    p.write_text("a,b")
    with p.open("rb") as fh:
        r = auth_client.post(
            "/api/files",
            files={"file": (p.name, fh)},
            data={"purpose": "training", "section_in": "8.5", "well_type": "J"},
        )
        assert r.status_code == 400


def test_full_pipeline(auth_client, db, inbox, tmp_dir):
    st = auth_client.get("/api/inbox/status").json()
    assert st["pending"] > 15 and st["wells"] == 10

    # --- impor massal (background task dijalankan sinkron oleh TestClient)
    r = auth_client.post("/api/inbox/scan")
    assert r.status_code == 200, r.text
    run = auth_client.get(f"/api/inbox/runs/{r.json()['id']}").json()
    assert run["status"] == "done", run
    assert run["counts"].get("rejected", 0) == 0, run["files"]
    assert not list(inbox.rglob("*.xls*")), "file harus dipindah dari inbox"
    assert list((tmp_dir / "processed").rglob("*.xls*"))
    # idempoten: memindai lagi tidak menemukan file
    assert auth_client.post("/api/inbox/scan").status_code == 400

    # --- gerbang kualitas
    q = auth_client.get("/api/quality").json()
    by = {(x["well"], x["section_in"]): x for x in q}
    new = [x for x in q if "BARU" in x["well"]]
    assert new and all(x["status"] == "C" for x in new)  # tanpa data aktual -> ditahan
    assert sum(x["status"] in ("A", "B") for x in q) >= 15
    for x in q:
        assert x["well_type"] in ("J", "S", "Horizontal")  # dari folder tipe
    # tinjauan tercatat
    target = new[0]
    rv = auth_client.post(
        f"/api/quality/{target['well_id']}/review",
        json={"decision": "fix", "reason": "minta data aktual ke client"},
    )
    assert rv.json()["status"] == "C"
    assert (
        auth_client.post(
            f"/api/quality/{target['well_id']}/review",
            json={"decision": "x", "reason": "not valid"},
        ).status_code
        == 400
    )
    xl = auth_client.get("/api/quality/report.xlsx")
    assert "Data quality" in openpyxl.load_workbook(io.BytesIO(xl.content)).sheetnames
    assert by  # noqa

    # --- bekukan dataset + blind test terkunci
    d = auth_client.post("/api/datasets/freeze").json()
    assert d["version"] == 1 and d["n_blind"] >= 1 and len(d["hash"]) == 64
    blind = auth_client.get("/api/datasets/blind").json()
    assert len(blind["wells"]) >= 1
    d2 = auth_client.post("/api/datasets/freeze").json()
    assert d2["version"] == 2 and d2["hash"] == d["hash"]  # data sama -> hash sama
    assert (
        auth_client.get("/api/datasets/blind").json()["wells"] == blind["wells"]
    )  # tetap terkunci

    # --- latih (sinkron agar deterministik)
    row = MLModel(algorithm="xgboost", status="queued")
    db.add(row)
    db.commit()
    train(db, row, "xgboost")
    assert row.status == "done" and row.active and row.dataset_id
    m = row.metrics
    assert m["dataset"]["wells_blind"] >= 1
    assert set(m["operations"]) >= {"pick_up", "slack_off"}
    op = m["operations"]["pick_up"]
    assert op["learning_curve"] and op["explain"]["features"] and op["strategy"]
    assert m["feature_selection"][0]["group"] == "base"
    assert 0 <= op["overall"]["within"]["ml"] <= 1
    assert (
        abs(op["overall"]["band_coverage"] - 0.8) < 0.05
    )  # pita P10–P90 dari residu OOF yang sama
    assert set(op["forecast_backtest"]["horizons"]) == {"300", "600", "1000"}
    assert "ML + bias" in op["forecast_backtest"]["horizons"]["300"]
    trained = set(m["dataset"]["train_combos"])
    assert trained
    # sumur blind tidak pernah punya prediksi out-of-fold
    oof_wells = {
        db.get(Well, wid).name
        for wid in db.scalars(
            select(Prediction.well_id).where(
                Prediction.kind == "oof", Prediction.model_id == row.id
            )
        )
    }
    assert oof_wells and not (oof_wells & set(blind["wells"]))

    # --- blind test sekali
    b = auth_client.post(f"/api/models/{row.id}/blind-test")
    assert b.status_code == 200, b.text
    assert set(b.json()["wells"]) == set(blind["wells"])
    assert all(0 <= v["band_coverage"] <= 1 for v in b.json()["operations"].values())
    assert auth_client.post(f"/api/models/{row.id}/blind-test").status_code == 400

    # --- laporan model
    assert auth_client.get(f"/api/models/{row.id}/report.pdf").content[:4] == b"%PDF"
    sheets = openpyxl.load_workbook(
        io.BytesIO(auth_client.get(f"/api/models/{row.id}/report.xlsx").content)
    ).sheetnames
    assert {"Summary", "Learning curve", "Feature importance", "Blind test", "Dataset"} <= set(
        sheets
    )

    # --- prediksi sumur baru + batas aman
    w_new = db.get(Well, target["well_id"])
    p = auth_client.post(f"/api/wells/{w_new.id}/predict")
    assert p.status_code == 200, p.text
    lim = auth_client.post(
        "/api/limits",
        json={
            "operation": "pick_up",
            "value": 50,
            "unit_system": "imperial",
            "section_in": w_new.section_in,
            "note": "uji",
        },
    )
    assert lim.status_code == 200, lim.text
    prof = auth_client.get(f"/api/wells/{w_new.id}/profile").json()
    assert prof["prediction"]["kind"] == "full"
    pu = prof["operations"]["pick_up"]
    assert pu["ml"]["lo"] and pu["ml"]["hi"]
    assert pu["limits"] and pu["limits"][0]["cross_ml"] is not None  # 50 klbf pasti terlampaui
    assert prof["quality"]["status"] == "C"

    x = auth_client.get(f"/api/wells/{w_new.id}/export.xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(x.content))
    # format template client "OUTPUT … Multiple T&D Road Map"
    assert wb.sheetnames[:4] == [
        "Summary Outputs",
        "Tripping Load Analysis - Graph",
        "Torque Analysis Off Btm",
        "Torque Analysis On Bottom",
    ]
    assert [n.split(" MW ")[0] for n in wb.sheetnames[4:]] == ["ROT", "SO", "PU"]
    sm = wb["Summary Outputs"]
    assert sm["C3"].value == "ML PREDICTION ANALYSIS SUMMARY REPORT"
    assert [sm.cell(16, c).value for c in (2, 3, 6, 9, 15)] == [
        "Bit Depth ",
        "Actual PU",
        "ML PU",
        "PU-MLPU",
        "TQ-MLTQ On",
    ]
    tr = wb["Tripping Load Analysis - Graph"]
    assert tr["Q2"].value == "MODELLED HOOKLOADS" and tr["Q3"].value == "Bit Depth"
    assert any(str(c.value).startswith("PU Hook Load ff=") for c in tr[3])
    assert "ML PREDICTION" in [c.value for c in tr[2]]
    tq = wb["Torque Analysis Off Btm"]
    assert tq["L3"].value == "MODELLED TORQUE" and tq["L4"].value == "Bit Depth"
    pu = wb[wb.sheetnames[-1]]
    assert pu["D5"].value == "Multipoint Torque and Drag Outputs"
    assert any("Trip out" in str(c.value) for r in pu.iter_rows() for c in r if c.value)
    assert "Multiple T&D Road Map" in x.headers["content-disposition"]
    assert auth_client.get(f"/api/wells/{w_new.id}/report.pdf").content[:4] == b"%PDF"

    # --- forecast N ft ke depan + penjelasan sebab-akibat + ekspor
    fc = auth_client.post(
        f"/api/wells/{w_new.id}/forecast", json={"distance_ft": 300, "bias_correction": True}
    )
    assert fc.status_code == 200, fc.text
    fc = fc.json()
    assert fc["end_depth"] - fc["start_depth"] == pytest.approx(300, abs=1)
    pu = fc["operations"]["pick_up"]
    assert len(pu["depth"]) == len(pu["ml"]) == len(pu["p10"]) == len(pu["p90"])
    assert pu["explanation"]["drivers"] and pu["explanation"]["sentence"].startswith("Pick up")
    assert pu["explanation"]["limit_crossings"]  # batas 50 klbf di atas
    assert fc["summary"]
    assert pu["backtest"] and 0 <= pu["backtest"]["within"] <= 1 and pu["tolerance"] == "8 klbf"
    assert fc["request"]["distance_ft"] == 300
    # trip in: garis peringatan set-down pada slack off (prediction - toleransi)
    so = fc["operations"]["slack_off"]
    assert fc["direction"] == "in" and so["alert"]["kind"] == "set_down"
    assert pu["alert"] is None
    # trip out: jendela di atas kedalaman awal, urut arah gerak bit, overpull alert pada pick up
    fo = auth_client.post(
        f"/api/wells/{w_new.id}/forecast",
        json={"distance_ft": 600, "direction": "out", "start_depth_ft": fc["end_depth"]},
    )
    assert fo.status_code == 200, fo.text
    fo = fo.json()
    assert fo["direction"] == "out" and fo["end_depth"] < fo["start_depth"]
    po = fo["operations"]["pick_up"]
    assert po["depth"][0] > po["depth"][-1]  # dari bawah ke atas
    assert po["alert"]["kind"] == "overpull" and po["backtest"] is None
    main = po["ml_corrected"] or po["ml"]
    assert po["alert"]["values"][0] == pytest.approx(main[0] + 8, abs=0.05)
    assert fo["operations"]["slack_off"]["alert"] is None
    xo = auth_client.get(
        f"/api/wells/{w_new.id}/export.xlsx?fc_distance_ft=600&fc_direction=out"
        f"&fc_start_depth_ft={fc['end_depth']}"
    )
    assert xo.status_code == 200, xo.text
    vals = [
        str(c.value)
        for row in openpyxl.load_workbook(io.BytesIO(xo.content))[
            "Tripping Load Analysis - Graph"
        ].iter_rows()
        for c in row
        if c.value
    ]
    assert any(v.startswith("TRIP OUT PREDICTION") for v in vals)
    assert any("Overpull alert" in v for v in vals)
    # ekspor Excel dengan prediction yang sedang tampil: tabel + garis ungu di grafik
    xe = auth_client.get(f"/api/wells/{w_new.id}/export.xlsx?fc_distance_ft=300&fc_bias=true")
    assert xe.status_code == 200, xe.text
    wbx = openpyxl.load_workbook(io.BytesIO(xe.content))
    for sh in ("Tripping Load Analysis - Graph", "Torque Analysis Off Btm"):
        vals = [str(c.value) for row in wbx[sh].iter_rows() for c in row if c.value]
        assert any(v.startswith("PREDICTION AHEAD") for v in vals), sh
        assert any("- Prediction P90" in v for v in vals), sh
    with zipfile.ZipFile(io.BytesIO(xe.content)) as z:
        charts = " ".join(z.read(n).decode() for n in z.namelist() if "charts/chart" in n)
    assert "4A3AA7" in charts.upper()

    # uji model pada sumur yang punya aktual: forecast dimulai sebelum aktual terakhir
    tw = next(
        w for w in db.scalars(select(Well)) if w.status == "ready" and w.name not in blind["wells"]
    )
    pr = auth_client.get(f"/api/wells/{tw.id}/profile").json()
    ad = sorted(pr["operations"]["pick_up"]["actual"]["depth"])
    mid = ad[len(ad) // 2]
    ck = auth_client.post(
        f"/api/wells/{tw.id}/forecast", json={"distance_ft": 600, "start_depth_ft": mid}
    ).json()
    c = ck["operations"]["pick_up"]["actual_check"]
    assert ck["bias_correction"] is True and c and c["n"] >= 1 and 0 <= c["ml_within"] <= 1
    assert any("before the last actual" in x for x in ck["warnings"])
    xf = auth_client.post(f"/api/wells/{w_new.id}/forecast.xlsx", json={"distance_ft": 300})
    assert {"Summary", "PU", "Explanation"} <= set(
        openpyxl.load_workbook(io.BytesIO(xf.content)).sheetnames
    )
    assert (
        auth_client.post(f"/api/wells/{w_new.id}/forecast", json={"distance_ft": -5}).status_code
        == 400
    )

    # --- evaluasi prediksi: data aktual sumur baru datang belakangan
    spec = well_specs(10, SEED)[9]
    data = build_well(section_spec(spec, w_new.section_in))
    from make_sample_data import _compute_plan

    _compute_plan(data, [0.2, 0.3, 0.4, 0.5])
    f = tmp_dir / f"P_{spec.name}_BHA_{w_new.section_in:g}in_TnD_aktual.xlsm"
    write_wellplan_file(data, f, include_actual=True)
    with f.open("rb") as fh:
        up = auth_client.post(
            "/api/files",
            files={"file": (f.name, fh)},
            data={
                "well_name": w_new.name,
                "purpose": "training",
                "section_in": str(w_new.section_in),
                "well_type": w_new.well_type,
            },
        )
    assert up.json()["status"] in ("ok", "warning"), up.json()
    evs = auth_client.get("/api/evaluations").json()
    assert evs and evs[0]["well_id"] == w_new.id
    assert "pick_up" in evs[0]["metrics"]["operations"]


def test_model_held_when_worse(db):
    from sqlalchemy import update

    db.execute(update(MLModel).values(active=False))
    db.commit()
    a = MLModel(algorithm="a", status="done", active=True, metrics={"skill": 0.7})
    b = MLModel(algorithm="b", status="running", metrics={"skill": 0.9})
    db.add_all([a, b])
    db.commit()
    _compare_and_activate(db, b)
    assert b.status == "held" and not b.active and "Held" in b.comparison["decision"]
    c = MLModel(algorithm="c", status="running", metrics={"skill": 0.6})
    db.add(c)
    db.commit()
    _compare_and_activate(db, c)
    db.commit()
    assert c.status == "done" and c.active
    db.refresh(a)
    assert not a.active
    for x in (a, b, c):
        db.delete(x)
    db.commit()
